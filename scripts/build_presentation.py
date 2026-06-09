from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG = RGBColor(0x1B, 0x28, 0x38)
ACCENT = RGBColor(0x66, 0xC0, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xC7, 0xD5, 0xE0)
GREEN = RGBColor(0x5C, 0xBA, 0x7D)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12), Inches(1.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12), Inches(0.8))
        p2 = box2.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = ACCENT


def add_bullets(slide, items, left=0.8, top=2.0, width=11.5, height=4.8, size=22):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = GRAY
        p.space_after = Pt(10)


def slide_title_only(title, subtitle):
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, title, subtitle)


def slide_content(title, bullets):
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, title)
    add_bullets(s, bullets)


def slide_two_columns(title, left_title, left_items, right_title, right_items):
    s = prs.slides.add_slide(blank)
    set_bg(s)
    add_title(s, title)
    lt = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.5), Inches(0.5))
    lt.text_frame.paragraphs[0].text = left_title
    lt.text_frame.paragraphs[0].font.size = Pt(24)
    lt.text_frame.paragraphs[0].font.bold = True
    lt.text_frame.paragraphs[0].font.color.rgb = ACCENT
    add_bullets(s, left_items, left=0.8, top=2.4, width=5.5, size=18)
    rt = s.shapes.add_textbox(Inches(7), Inches(1.8), Inches(5.5), Inches(0.5))
    rt.text_frame.paragraphs[0].text = right_title
    rt.text_frame.paragraphs[0].font.size = Pt(24)
    rt.text_frame.paragraphs[0].font.bold = True
    rt.text_frame.paragraphs[0].font.color.rgb = ACCENT
    add_bullets(s, right_items, left=7, top=2.4, width=5.5, size=18)


slide_title_only(
    "CtrlGame",
    "Веб-приложение цифрового магазина игр\nИркутский государственный университет · 2026",
)

slide_content("О проекте", [
    "CtrlGame (ControlGame) — учебный fullstack-проект",
    "Цифровой магазин игр в стиле Steam",
    "Каталог, корзина, покупки с баланса, библиотека, отзывы",
    "Роли: пользователь и администратор",
    "Репозиторий: github.com/quasponn/ctrlgame",
])

slide_content("Цель и задачи", [
    "Цель: веб-приложение с REST API, БД и современным UI",
    "Спроектировать базу данных (SQLite, 11 таблиц)",
    "Реализовать backend на Node.js + Express",
    "Разработать frontend на React",
    "Авторизация JWT, роли, корзина и покупки",
    "Автотесты и CI на GitHub Actions",
])

slide_content("Технологический стек", [
    "Frontend: React 19, Vite, React Router",
    "Backend: Node.js, Express 5",
    "База данных: SQLite (файл db.sqlite)",
    "Безопасность: JWT, bcrypt, Helmet, rate limit",
    "CI: GitHub Actions (тесты + сборка)",
    "Оформление: тёмная тема в духе Steam",
])

slide_content("Архитектура", [
    "Клиент–сервер: браузер ↔ Express ↔ SQLite",
    "Один сервер на порту 3001 — API + собранный frontend",
    "Запросы к /api/... → JSON",
    "Frontend: pages, components, AuthContext, api/client.js",
    "Backend: routes/, services/, middleware/, db.js",
])

slide_two_columns(
    "Команда проекта",
    "Артём — Backend",
    [
        "База данных (db.js, seed.js)",
        "REST API (routes/)",
        "JWT + refresh-токены",
        "Каталог, корзина, checkout",
        "Автотесты (10 шт.)",
        "GitHub Actions (CI)",
    ],
    "Галсан — Frontend",
    [
        "React, Vite, Router",
        "Страницы: магазин, игра, библиотека",
        "Компоненты UI, корзина, акции",
        "AuthContext, ToastContext",
        "api/client.js",
        "Стили (Steam-тема)",
    ],
)

slide_content("База данных (11 таблиц)", [
    "users — аккаунты, баланс, роль",
    "games — каталог, цена, discount_percent",
    "genres, game_genres — жанры (M:N)",
    "cart, library, wishlist — корзина, покупки, желаемое",
    "reviews, game_screenshots, game_patches",
    "refresh_tokens — сессии",
    "Схема: docs/database-schema.mmd → mermaid.live",
])

slide_content("REST API", [
    "Публично: каталог, акции, игра, отзывы, login/register",
    "С JWT: корзина, checkout, библиотека, wishlist, профиль",
    "Только admin: создание и редактирование игр",
    "Каталог: поиск, жанр, пагинация на сервере",
    "Checkout: цена со скидкой, проверка баланса → library",
])

slide_content("Функционал для пользователя", [
    "Магазин: карусель, акции, поиск, фильтры, пагинация",
    "Страница игры: галерея, отзывы, патчи, рейтинг",
    "Корзина и покупка с подтверждением",
    "Библиотека купленных игр",
    "Профиль: данные, смена пароля, пополнение баланса",
    "Список желаемого (wishlist)",
])

slide_content("Администратор и безопасность", [
    "Admin: кнопка «+ Игра», CRUD игр и жанров",
    "POST /api/games защищён middleware requireAdmin",
    "Пароли: bcrypt-хеш, не plain text",
    "Access-токен (~1 ч) + refresh (~30 дней)",
    "Rate limit на login/register",
    "CORS, секреты в .env",
])

slide_content("Тестирование и CI", [
    "10 автотестов в backend/test/",
    "Каталог, пагинация, валидация, login/refresh",
    "Запрет создания игр обычным user (403)",
    "Запуск: cd backend → npm test",
    "GitHub Actions: seed → test → build frontend",
])

slide_content("Запуск проекта", [
    "PowerShell из корня проекта:",
    "  .\\scripts\\setup.ps1   — первый раз",
    "  .\\scripts\\start.ps1    — сборка + запуск",
    "Сайт: http://localhost:3001",
    "Логины: admin/admin123, user/user123",
])

slide_content("Демонстрация (сценарий)", [
    "1. Каталог и акции без входа",
    "2. Вход user → wishlist → корзина",
    "3. Пополнение баланса → покупка",
    "4. Библиотека — купленная игра",
    "5. Отзыв на странице игры",
    "6. Вход admin → добавление игры",
])

slide_title_only(
    "Спасибо за внимание!",
    "Вопросы?\nCtrlGame · http://localhost:3001",
)

out = r"D:\game-library\docs\CtrlGame-presentation.pptx"
prs.save(out)
print(f"Saved: {out}")
